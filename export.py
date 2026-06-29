"""
Export Module
Export statistics and tables to Excel, CSV, and Word-ready formats
"""

import pandas as pd
import numpy as np
from pathlib import Path
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

import config


# ============================================================================
# EXCEL EXPORT
# ============================================================================

class ExcelExporter:
    """Export statistics to formatted Excel workbooks."""
    
    def __init__(self, filename):
        """
        Initialize Excel exporter.
        
        Parameters
        ----------
        filename : str
            Output filename
        """
        self.filepath = config.TABLES_DIR / f"{filename}.xlsx"
        self.wb = Workbook()
        self.wb.remove(self.wb.active)
    
    def add_styled_dataframe(self, df, sheet_name, title=None, 
                            freeze_panes=(1, 0)):
        """
        Add DataFrame to worksheet with formatting.
        
        Parameters
        ----------
        df : pd.DataFrame
            Data to add
        sheet_name : str
            Name of worksheet
        title : str, optional
            Sheet title
        freeze_panes : tuple
            Position to freeze panes (row, col)
        """
        ws = self.wb.create_sheet(sheet_name)
        
        # Add title
        if title:
            ws['A1'] = title
            ws['A1'].font = Font(size=14, bold=True)
            ws.merge_cells('A1:Z1')
            start_row = 2
        else:
            start_row = 0
        
        # Add DataFrame
        for r_idx, row in enumerate(df.itertuples(index=False, name=None),
                                    start=start_row + 1):
            for c_idx, value in enumerate(row, start=1):
                cell = ws.cell(row=r_idx, column=c_idx, value=value)
                
                # Format numbers
                if isinstance(value, (int, float)):
                    if c_idx == 1 or isinstance(value, int):
                        cell.number_format = '0'
                    else:
                        cell.number_format = '0.0000'
        
        # Add headers
        for c_idx, column in enumerate(df.columns, start=1):
            cell = ws.cell(row=start_row + 1, column=c_idx, value=column)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill(start_color="366092", end_color="366092",
                                   fill_type="solid")
            cell.alignment = Alignment(horizontal="center", vertical="center")
        
        # Auto-adjust column widths
        for c_idx, column in enumerate(df.columns, start=1):
            max_length = max(
                len(str(column)),
                max(len(str(row[c_idx-1])) for row in df.itertuples(index=False))
            )
            ws.column_dimensions[get_column_letter(c_idx)].width = min(max_length + 2, 50)
        
        # Freeze panes
        if freeze_panes:
            ws.freeze_panes = f"{get_column_letter(freeze_panes[1]+1)}{freeze_panes[0]+1}"
        
        return ws
    
    def save(self):
        """Save workbook."""
        self.wb.save(self.filepath)
        if config.VERBOSE >= 1:
            print(f"  ✓ Saved {self.filepath.name}")


# ============================================================================
# SUMMARY STATISTICS EXPORT
# ============================================================================

def export_summary_statistics(all_results):
    """
    Export comprehensive summary statistics to Excel.
    
    Parameters
    ----------
    all_results : dict
        Dictionary of {complex: {descriptor: statistics}}
    """
    exporter = ExcelExporter("summary_statistics")
    
    # Overall summary by complex
    for complex_name, descriptors in all_results.items():
        data = []
        for desc_name, stats in descriptors.items():
            data.append({
                'Descriptor': desc_name,
                'Mean': f"{stats['descriptive']['mean']:.6f}",
                'Median': f"{stats['descriptive']['median']:.6f}",
                'Std Dev': f"{stats['descriptive']['std']:.6f}",
                'SEM': f"{stats['descriptive']['sem']:.6f}",
                'CV (%)': f"{stats['descriptive']['cv']:.2f}",
                'Min': f"{stats['descriptive']['min']:.6f}",
                'Max': f"{stats['descriptive']['max']:.6f}",
                'IQR': f"{stats['descriptive']['iqr']:.6f}",
            })
        
        df = pd.DataFrame(data)
        exporter.add_styled_dataframe(df, complex_name, 
                                     title=f"{complex_name} - Descriptive Statistics")
    
    exporter.save()


def export_bootstrap_results(all_results):
    """
    Export bootstrap confidence intervals.
    
    Parameters
    ----------
    all_results : dict
        Dictionary of {complex: {descriptor: statistics}}
    """
    exporter = ExcelExporter("bootstrap_results")
    
    for complex_name, descriptors in all_results.items():
        data = []
        for desc_name, stats in descriptors.items():
            bootstrap = stats['bootstrap']
            data.append({
                'Descriptor': desc_name,
                'Mean': f"{bootstrap['mean']:.6f}",
                'Bootstrap Mean': f"{bootstrap['bootstrap_mean']:.6f}",
                'Bootstrap SD': f"{bootstrap['bootstrap_std']:.6f}",
                'CI Lower (95%)': f"{bootstrap['ci_lower_95']:.6f}",
                'CI Upper (95%)': f"{bootstrap['ci_upper_95']:.6f}",
                'CI Width': f"{bootstrap['ci_width']:.6f}",
                'Method': bootstrap['method'].upper(),
            })
        
        df = pd.DataFrame(data)
        exporter.add_styled_dataframe(df, complex_name,
                                     title=f"{complex_name} - Bootstrap Confidence Intervals")
    
    exporter.save()


def export_effective_sample_size(all_results):
    """
    Export effective sample size calculations.
    
    Parameters
    ----------
    all_results : dict
        Dictionary of {complex: {descriptor: statistics}}
    """
    exporter = ExcelExporter("effective_sample_size")
    
    for complex_name, descriptors in all_results.items():
        data = []
        for desc_name, stats in descriptors.items():
            ess = stats['ess']
            data.append({
                'Descriptor': desc_name,
                'N Total': int(ess['n_total']),
                'Tau Integrated': f"{ess['tau_integrated']:.4f}",
                'N Effective': f"{ess['n_effective']:.1f}",
                'Efficiency (%)': f"{ess['efficiency']:.2f}",
                'SEM (uncorrected)': f"{ess['sem_uncorrected']:.6f}",
                'SEM (corrected)': f"{ess['sem_corrected']:.6f}",
                'CI Width (uncorrected)': f"{ess['ci_width_95_uncorrected']:.6f}",
                'CI Width (corrected)': f"{ess['ci_width_95_corrected']:.6f}",
            })
        
        df = pd.DataFrame(data)
        exporter.add_styled_dataframe(df, complex_name,
                                     title=f"{complex_name} - Effective Sample Size")
    
    exporter.save()


def export_trend_analysis(all_results):
    """
    Export linear trend analysis results.
    
    Parameters
    ----------
    all_results : dict
        Dictionary of {complex: {descriptor: statistics}}
    """
    exporter = ExcelExporter("trend_analysis")
    
    for complex_name, descriptors in all_results.items():
        data = []
        for desc_name, stats in descriptors.items():
            trend = stats['trend']
            converged = "No" if trend['significant'] else "Yes"
            data.append({
                'Descriptor': desc_name,
                'Slope': f"{trend['slope']:.8f}",
                'Intercept': f"{trend['intercept']:.6f}",
                'R²': f"{trend['r_squared']:.6f}",
                'R value': f"{trend['r_value']:.6f}",
                'P-value': f"{trend['p_value']:.4e}",
                'Std Error': f"{trend['std_err']:.8f}",
                'Converged (No trend)': converged,
            })
        
        df = pd.DataFrame(data)
        exporter.add_styled_dataframe(df, complex_name,
                                     title=f"{complex_name} - Trend Analysis")
    
    exporter.save()


def export_convergence_analysis(all_results):
    """
    Export convergence test results.
    
    Parameters
    ----------
    all_results : dict
        Dictionary of {complex: {descriptor: statistics}}
    """
    exporter = ExcelExporter("convergence_analysis")
    
    for complex_name, descriptors in all_results.items():
        # Window statistics
        data_windows = []
        data_tests = []
        
        for desc_name, stats in descriptors.items():
            convergence = stats.get('convergence', {})
            
            window_stats = convergence.get('window_statistics', {})
            for window_name, win_stats in window_stats.items():
                data_windows.append({
                    'Descriptor': desc_name,
                    'Window': window_name,
                    'N': int(win_stats['n']),
                    'Mean': f"{win_stats['mean']:.6f}",
                    'Std': f"{win_stats['std']:.6f}",
                    'SEM': f"{win_stats['sem']:.6f}",
                    'Min': f"{win_stats['min']:.6f}",
                    'Max': f"{win_stats['max']:.6f}",
                })
            
            # Statistical tests
            tests = convergence.get('statistical_tests', {})
            for test_name, test_result in tests.items():
                converged = "Yes" if test_result['converged'] else "No"
                data_tests.append({
                    'Descriptor': desc_name,
                    'Test': test_name,
                    't-statistic': f"{test_result['t_statistic']:.6f}",
                    'P-value': f"{test_result['p_value']:.4e}",
                    'Converged': converged,
                })
        
        if data_windows:
            df = pd.DataFrame(data_windows)
            exporter.add_styled_dataframe(df, f"{complex_name}_windows",
                                         title=f"{complex_name} - Convergence Window Statistics")
        
        if data_tests:
            df = pd.DataFrame(data_tests)
            exporter.add_styled_dataframe(df, f"{complex_name}_tests",
                                         title=f"{complex_name} - Convergence Tests")
    
    exporter.save()


# ============================================================================
# PUBLICATION TABLE EXPORT
# ============================================================================

def export_publication_table(all_results, output_format='excel'):
    """
    Create publication-ready summary table for all complexes.
    
    Parameters
    ----------
    all_results : dict
        Dictionary of {complex: {descriptor: statistics}}
    output_format : str
        Format: 'excel', 'csv', or 'both'
    """
    data = []
    
    for complex_name, descriptors in all_results.items():
        for desc_name, stats in descriptors.items():
            desc_stat = stats['descriptive']
            bootstrap = stats['bootstrap']
            ess = stats['ess']
            trend = stats['trend']
            
            row = {
                'Complex': complex_name,
                'Descriptor': desc_name,
                'Mean ± SD': f"{desc_stat['mean']:.4f} ± {desc_stat['std']:.4f}",
                '95% CI': f"[{bootstrap['ci_lower_95']:.4f}, {bootstrap['ci_upper_95']:.4f}]",
                'CV (%)': f"{desc_stat['cv']:.2f}",
                'Neff': f"{ess['n_effective']:.0f}",
                'Trend Slope': f"{trend['slope']:.2e}",
                'R²': f"{trend['r_squared']:.4f}",
                'Converged': "Yes" if not trend['significant'] else "No",
            }
            data.append(row)
    
    df_pub = pd.DataFrame(data)
    
    if output_format in ['excel', 'both']:
        filepath_excel = config.TABLES_DIR / "publication_table.xlsx"
        df_pub.to_excel(filepath_excel, index=False, sheet_name="Results")
        if config.VERBOSE >= 1:
            print(f"  ✓ Saved {filepath_excel.name}")
    
    if output_format in ['csv', 'both']:
        filepath_csv = config.TABLES_DIR / "publication_table.csv"
        df_pub.to_csv(filepath_csv, index=False)
        if config.VERBOSE >= 1:
            print(f"  ✓ Saved {filepath_csv.name}")
    
    return df_pub


# ============================================================================
# MARKDOWN TABLE EXPORT (for Word, presentations, etc.)
# ============================================================================

def export_markdown_table(all_results, filename="tables_markdown"):
    """
    Export tables in Markdown format (easy to copy to Word).
    
    Parameters
    ----------
    all_results : dict
        Dictionary of {complex: {descriptor: statistics}}
    filename : str
        Output filename without extension
    """
    filepath = config.TABLES_DIR / f"{filename}.md"
    
    with open(filepath, 'w') as f:
        f.write("# MD Simulation Analysis Results\n\n")
        
        # Publication table
        f.write("## Summary Table\n\n")
        f.write("| Complex | Descriptor | Mean ± SD | 95% CI | CV (%) | Neff | Converged |\n")
        f.write("|---------|------------|-----------|--------|--------|------|----------|\n")
        
        for complex_name, descriptors in all_results.items():
            for desc_name, stats in descriptors.items():
                desc_stat = stats['descriptive']
                bootstrap = stats['bootstrap']
                ess = stats['ess']
                trend = stats['trend']
                
                converged = "Yes" if not trend['significant'] else "No"
                f.write(
                    f"| {complex_name} | {desc_name} | "
                    f"{desc_stat['mean']:.4f} ± {desc_stat['std']:.4f} | "
                    f"[{bootstrap['ci_lower_95']:.4f}, {bootstrap['ci_upper_95']:.4f}] | "
                    f"{desc_stat['cv']:.2f} | {ess['n_effective']:.0f} | {converged} |\n"
                )
        
        # Detailed tables per complex
        for complex_name, descriptors in all_results.items():
            f.write(f"\n## {complex_name} - Detailed Statistics\n\n")
            f.write("| Descriptor | Mean | Median | SD | CV (%) | Min | Max |\n")
            f.write("|------------|------|--------|----|---------|----|-----|\n")
            
            for desc_name, stats in descriptors.items():
                desc_stat = stats['descriptive']
                f.write(
                    f"| {desc_name} | {desc_stat['mean']:.6f} | {desc_stat['median']:.6f} | "
                    f"{desc_stat['std']:.6f} | {desc_stat['cv']:.2f} | "
                    f"{desc_stat['min']:.6f} | {desc_stat['max']:.6f} |\n"
                )
    
    if config.VERBOSE >= 1:
        print(f"  ✓ Saved {filepath.name}")


# ============================================================================
# CSV EXPORT
# ============================================================================

def export_block_averaged_data(all_results):
    """
    Export block-averaged data for external analysis.
    
    Parameters
    ----------
    all_results : dict
        Dictionary of {complex: {descriptor: data}}
    """
    for complex_name, descriptors in all_results.items():
        data = []
        
        for desc_name, result in descriptors.items():
            block_time = result.get('block_time')
            block_means = result.get('block_means')
            
            if block_time is not None and block_means is not None:
                for t, val in zip(block_time, block_means):
                    data.append({
                        'Time_ns': f"{t:.3f}",
                        f'{desc_name}': f"{val:.6f}",
                    })
        
        if data:
            df = pd.DataFrame(data)
            filepath = config.DATA_DIR / f"{complex_name}_block_averaged.csv"
            df.to_csv(filepath, index=False)
            if config.VERBOSE >= 1:
                print(f"  ✓ Saved {filepath.name}")
